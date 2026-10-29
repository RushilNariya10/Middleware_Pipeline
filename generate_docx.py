"""
Generate 24IT055-PR-7.docx — Practical 7: Authentication & Middleware Pipeline
"""
import os, sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

BASE = r"d:\F\SEM 5\AWDF\P7"
OUT  = os.path.join(BASE, "24IT055-PR-7.docx")

doc = Document()

# ── Page margins ──
for section in doc.sections:
    section.top_margin    = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin   = Cm(2.5)
    section.right_margin  = Cm(2.5)

style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)
style.paragraph_format.space_after = Pt(4)

def add_heading_styled(text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x0D, 0x6E, 0xFD)
    return h

def add_code_block(code, language="javascript"):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(code)
    run.font.name = 'Consolas'
    run.font.size = Pt(9)
    # Light grey shading
    shading = run._element.get_or_add_rPr()
    shade_elm = shading.makeelement(qn('w:shd'), {
        qn('w:val'): 'clear',
        qn('w:color'): 'auto',
        qn('w:fill'): 'F5F5F5'
    })
    shading.append(shade_elm)

def add_table(headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Light Grid Accent 1'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0]
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        cell.text = h
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(9)
    for row_data in rows:
        row = table.add_row()
        for i, val in enumerate(row_data):
            cell = row.cells[i]
            cell.text = str(val)
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)

# ══════════════════════════════════════════════════
# HEADER
# ══════════════════════════════════════════════════

# University logos
logo_table = doc.add_table(rows=1, cols=3)
logo_table.alignment = WD_TABLE_ALIGNMENT.CENTER
for cell in logo_table.rows[0].cells:
    cell.width = Inches(2)

cspit_logo = os.path.join(BASE, "task-api", "assets", "cspit_logo.png")
depstar_logo = os.path.join(BASE, "task-api", "assets", "depstar_logo.png")

if os.path.exists(cspit_logo):
    p = logo_table.rows[0].cells[0].paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run()
    r.add_picture(cspit_logo, width=Inches(1.2))

center_cell = logo_table.rows[0].cells[1]
cp = center_cell.paragraphs[0]
cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = cp.add_run("CHAROTAR UNIVERSITY OF\nSCIENCE AND TECHNOLOGY")
r.bold = True
r.font.size = Pt(10)

if os.path.exists(depstar_logo):
    p = logo_table.rows[0].cells[2].paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run()
    r.add_picture(depstar_logo, width=Inches(1.2))

doc.add_paragraph()  # spacer

# University details
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("FACULTY OF TECHNOLOGY AND ENGINEERING\n")
r.bold = True
r.font.size = Pt(11)
r = p.add_run("B. Tech (IT/CE/CSE/AIML)\n")
r.font.size = Pt(10)
r = p.add_run("Subject: Advanced Web Development Frameworks (ITUE301) | Semester: 5th")
r.font.size = Pt(10)

doc.add_paragraph()  # spacer

# Practical title
h = doc.add_heading("Practical 7: Authentication and Middleware Pipeline", level=1)
h.alignment = WD_ALIGN_PARAGRAPH.CENTER
for run in h.runs:
    run.font.color.rgb = RGBColor(0x0D, 0x6E, 0xFD)

# Student details
details = [
    ("Student Name", "Rushil Nariya"),
    ("Student ID", "24IT055"),
    ("GitHub Repository", "https://github.com/RushilNariya10/Task_Manager_API"),
    ("Tech Stack", "JavaScript (ES6+), Express.js, Node.js, MongoDB, Mongoose, bcryptjs, JSON Web Tokens, React, Vite"),
    ("Date", "29 October 2026"),
]
for label, value in details:
    p = doc.add_paragraph()
    r = p.add_run(f"{label}: ")
    r.bold = True
    r.font.size = Pt(10)
    r = p.add_run(value)
    r.font.size = Pt(10)

doc.add_page_break()

# ══════════════════════════════════════════════════
# 1. OBJECTIVE
# ══════════════════════════════════════════════════

add_heading_styled("1. Objective", level=1)
doc.add_paragraph(
    "Implement a complete authentication and authorization layer for the existing Task Manager REST API. "
    "This involves adding user registration with secure password hashing, JWT-based login, "
    "an authentication middleware that protects all task routes, and a server-side validation middleware "
    "that sanitises and validates incoming data before it reaches the database. The frontend is extended "
    "with login/register forms, token-based API calls, and automatic 401 handling."
)

# ══════════════════════════════════════════════════
# 2. PROBLEM STATEMENT
# ══════════════════════════════════════════════════

add_heading_styled("2. Problem Statement", level=1)
problems = [
    "Add user registration and login functionality to the Task Management application.",
    "Hash passwords using bcrypt before saving the User document to MongoDB.",
    "Implement JWT token generation on successful login, with a 1-hour expiry.",
    "Protect all task routes using an authentication middleware that verifies the JWT.",
    "Add server-side input validation that rejects malformed requests before they reach the database.",
]
for prob in problems:
    doc.add_paragraph(prob, style='List Bullet')

doc.add_paragraph()
doc.add_paragraph("Supplementary Problems:", style='List Bullet')
supp = [
    "GET /api/auth/me — a protected endpoint returning the current user's profile.",
    "Frontend: Register and Login pages, localStorage token management, Bearer header on every request.",
    "401 handling: clear token and redirect to login when any request returns 401.",
    "Logout button that clears the token and navigates to the login page.",
    "Hide the Tasks UI when the user is not authenticated.",
]
for s in supp:
    doc.add_paragraph(s, style='List Bullet 2')

doc.add_page_break()

# ══════════════════════════════════════════════════
# 3. ARCHITECTURE / FLOW DIAGRAM
# ══════════════════════════════════════════════════

add_heading_styled("3. Architecture: Authentication & Middleware Pipeline", level=1)

doc.add_paragraph(
    "The diagram below illustrates the request lifecycle. Every incoming request first passes through "
    "global middleware (CORS, JSON parsing). Protected routes then go through the JWT authentication "
    "middleware, followed by the validation middleware, before reaching the controller logic."
)

add_code_block(
"""Client (React / Postman)
    |  HTTP Request with Authorization: Bearer <token>
    v
+-----------------------------------------------------------+
|                   Express Application                      |
|                                                           |
|  [1] CORS + JSON Body Parser (global)                     |
|       |                                                   |
|  [2] Auth Middleware (middleware/auth.js)                  |
|       - Extracts Bearer token from Authorization header   |
|       - jwt.verify() with JWT_SECRET                      |
|       - Rejects 401 if missing, tampered, or expired      |
|       - Attaches decoded { id } to req.user               |
|       |                                                   |
|  [3] Validation Middleware (middleware/validate.js)        |
|       - Sanitises input (trim, strip HTML tags)           |
|       - Checks required fields, types, lengths            |
|       - Returns 400 with errors array if invalid          |
|       |                                                   |
|  [4] Route Controller                                     |
|       - Performs CRUD via Mongoose model methods           |
|       |                                                   |
|  [5] Global Error Handler                                 |
|       - Catches ValidationError, CastError, 11000         |
|       - Returns structured JSON, never raw stack traces   |
+-----------------------------------------------------------+
    |
    v
+-----------------------------------------------------------+
| Mongoose ODM                                              |
|  - User Schema: email (unique, lowercase), password hash  |
|  - Task Schema: title (required), priority (enum), etc.   |
|  - Pre-save hook: bcrypt.hash(password, 10)               |
|  - toJSON transform: strips password from responses       |
+-----------------------------------------------------------+
    |
    v
+-----------------------------------------------------------+
| MongoDB Atlas                                             |
|  - users collection   (email, password hash, timestamps)  |
|  - tasks collection   (title, description, priority, ...) |
+-----------------------------------------------------------+"""
)

doc.add_page_break()

# ══════════════════════════════════════════════════
# 4. STEP-BY-STEP IMPLEMENTATION
# ══════════════════════════════════════════════════

add_heading_styled("4. Step-by-Step Implementation", level=1)

# 4.1 User Model
add_heading_styled("4.1 User Model (models/User.js)", level=2)
doc.add_paragraph(
    "The User model defines an email field (required, unique, lowercase, trimmed) and a password field. "
    "A pre-save hook hashes the password with bcrypt (10 salt rounds) before storing it. "
    "The toJSON method ensures the password hash is never included in API responses."
)
add_code_block(
"""const mongoose = require('mongoose');
const bcrypt = require('bcryptjs');

const userSchema = new mongoose.Schema({
  email: {
    type: String,
    required: [true, 'Email is required'],
    unique: true,
    lowercase: true,
    trim: true
  },
  password: {
    type: String,
    required: [true, 'Password is required']
  }
}, { timestamps: true });

// Hash password before saving
userSchema.pre('save', async function () {
  if (!this.isModified('password')) return;
  this.password = await bcrypt.hash(this.password, 10);
});

// Compare candidate password with stored hash
userSchema.methods.comparePassword = async function (candidatePassword) {
  return bcrypt.compare(candidatePassword, this.password);
};

// Strip password from JSON output
userSchema.methods.toJSON = function () {
  const obj = this.toObject();
  delete obj.password;
  return obj;
};

module.exports = mongoose.model('User', userSchema);"""
)

# 4.2 Auth Routes
add_heading_styled("4.2 Auth Routes (routes/auth.js)", level=2)
doc.add_paragraph(
    "Three endpoints handle authentication: POST /api/auth/register creates a user after validating "
    "input and checking for duplicates (409). POST /api/auth/login finds the user, compares the password "
    "hash, and returns a signed JWT on success. Both use the same generic 'Invalid credentials' message "
    "for wrong email or password to prevent user enumeration. GET /api/auth/me is a protected endpoint "
    "that returns the authenticated user's details."
)
add_code_block(
"""// POST /api/auth/register
router.post('/register', validateRegister, async (req, res, next) => {
  const { email, password } = req.body;
  const existingUser = await User.findOne({ email: email.toLowerCase() });
  if (existingUser) {
    return res.status(409).json({ message: 'An account with this email already exists.' });
  }
  const user = await User.create({ email, password });
  res.status(201).json({ message: 'User registered successfully.', user });
});

// POST /api/auth/login
router.post('/login', validateLogin, async (req, res, next) => {
  const user = await User.findOne({ email: email.toLowerCase() });
  if (!user) return res.status(401).json({ message: 'Invalid credentials.' });

  const isMatch = await user.comparePassword(password);
  if (!isMatch) return res.status(401).json({ message: 'Invalid credentials.' });

  const token = jwt.sign({ id: user._id }, process.env.JWT_SECRET, { expiresIn: '1h' });
  res.status(200).json({ message: 'Login successful.', token, user });
});

// GET /api/auth/me (protected)
router.get('/me', authMiddleware, async (req, res, next) => {
  const user = await User.findById(req.user.id);
  res.status(200).json({ user });
});"""
)

# 4.3 Auth Middleware
add_heading_styled("4.3 Authentication Middleware (middleware/auth.js)", level=2)
doc.add_paragraph(
    "This middleware reads the Authorization header, verifies the Bearer token using jwt.verify(), "
    "and attaches the decoded payload to req.user. It wraps the verification in try/catch so that "
    "invalid, tampered, or expired tokens return 401 without crashing the server. It differentiates "
    "between TokenExpiredError and other JWT errors for clearer error messages."
)
add_code_block(
"""function authMiddleware(req, res, next) {
  const header = req.headers.authorization;
  if (!header || !header.startsWith('Bearer ')) {
    return res.status(401).json({ message: 'Access denied. No token provided.' });
  }

  const token = header.split(' ')[1];
  try {
    const decoded = jwt.verify(token, process.env.JWT_SECRET);
    req.user = decoded;
    next();
  } catch (err) {
    if (err.name === 'TokenExpiredError') {
      return res.status(401).json({ message: 'Token has expired. Please log in again.' });
    }
    return res.status(401).json({ message: 'Invalid token. Authentication failed.' });
  }
}"""
)

# 4.4 Validation Middleware
add_heading_styled("4.4 Validation Middleware (middleware/validate.js)", level=2)
doc.add_paragraph(
    "Three validation functions run before controllers: validateRegister checks email format and "
    "minimum password length. validateLogin ensures both fields are present. validateTask requires "
    "a title on POST, enforces length limits, validates priority enum values, and sanitises all "
    "string inputs by stripping HTML tags and trimming whitespace. All return 400 with a structured "
    "errors array on failure."
)
add_code_block(
"""function sanitize(value) {
  if (typeof value !== 'string') return value;
  return value.replace(/<[^>]*>/g, '').trim();
}

function validateTask(req, res, next) {
  const errors = [];
  if (title !== undefined) req.body.title = sanitize(title);
  if (req.method === 'POST' && !req.body.title) {
    errors.push('Task title is required.');
  }
  if (req.body.title && req.body.title.length > 200) {
    errors.push('Task title must not exceed 200 characters.');
  }
  if (errors.length > 0) {
    return res.status(400).json({ message: 'Validation Error', errors });
  }
  next();
}"""
)

# 4.5 Protected Task Routes
add_heading_styled("4.5 Protecting Task Routes (server.js)", level=2)
doc.add_paragraph(
    "All five task routes now include authMiddleware as the first handler. Create and update routes "
    "additionally include validateTask after authMiddleware. The pipeline order is always: "
    "Auth → Validation → Controller."
)
add_code_block(
"""// Pipeline: Auth → Validate → Controller
app.post('/tasks', authMiddleware, validateTask, async (req, res, next) => { ... });
app.put('/tasks/:id', authMiddleware, validateTask, async (req, res, next) => { ... });

// Read and delete only need auth
app.get('/tasks', authMiddleware, async (req, res, next) => { ... });
app.delete('/tasks/:id', authMiddleware, async (req, res, next) => { ... });"""
)

# 4.6 Frontend
add_heading_styled("4.6 Frontend Authentication (React)", level=2)
doc.add_paragraph(
    "An AuthContext (src/context/AuthContext.jsx) manages the authentication state across the app. "
    "It stores the JWT token in localStorage and provides login(), register(), logout(), and authFetch() "
    "functions. The authFetch wrapper automatically injects the Bearer header and handles 401 responses "
    "by clearing the token and redirecting to login."
)
doc.add_paragraph(
    "New pages Login.jsx and Register.jsx provide forms for authentication. Tasks.jsx is a full CRUD "
    "interface for managing tasks. A ProtectedRoute component wraps the Tasks page, redirecting "
    "unauthenticated users to /login. The Navbar conditionally shows the Tasks link and Logout button "
    "only when authenticated."
)
add_code_block(
"""// AuthContext — authFetch wrapper (handles 401 automatically)
const authFetch = async (url, options = {}) => {
  const headers = { ...options.headers, 'Authorization': `Bearer ${token}` };
  const res = await fetch(`${API_BASE}${url}`, { ...options, headers });
  if (res.status === 401) {
    logout();  // Clear token, redirect to login
    throw new Error('Session expired. Please log in again.');
  }
  return res;
};"""
)

doc.add_page_break()

# ══════════════════════════════════════════════════
# 5. TESTING & OUTPUT
# ══════════════════════════════════════════════════

add_heading_styled("5. Testing Evidence & Output", level=1)

doc.add_paragraph(
    "All 18 automated tests pass, covering: registration, duplicate rejection, login, wrong password, "
    "JWT verification, missing/malformed/tampered/expired tokens, protected CRUD operations, "
    "validation errors, and bcrypt hash verification in the database."
)

doc.add_paragraph(
    "The automated test suite (test-auth.js) is run with: npm run test:auth"
)

add_heading_styled("5.1 Test Results Summary", level=2)

add_table(
    ["#", "Test Case", "Expected", "Result"],
    [
        ["1",  "Register with valid data",          "201 Created",     "✔ PASS"],
        ["2",  "Duplicate email registration",       "409 Conflict",    "✔ PASS"],
        ["3",  "Register missing email",             "400 Bad Request", "✔ PASS"],
        ["4",  "Register short password (<6 chars)",  "400 Bad Request", "✔ PASS"],
        ["5",  "Login with valid credentials",       "200 OK + token",  "✔ PASS"],
        ["6",  "Login wrong password",               "401 Unauthorized","✔ PASS"],
        ["7",  "Login non-existent email",           "401 Unauthorized","✔ PASS"],
        ["8",  "GET /api/auth/me with valid token",  "200 OK + user",   "✔ PASS"],
        ["9",  "GET /tasks with no token",           "401 Unauthorized","✔ PASS"],
        ["10", "GET /tasks malformed header",        "401 Unauthorized","✔ PASS"],
        ["11", "GET /tasks tampered token",          "401 Unauthorized","✔ PASS"],
        ["12", "GET /tasks expired token (1s TTL)",  "401 Unauthorized","✔ PASS"],
        ["13", "POST /tasks create with auth",       "201 Created",     "✔ PASS"],
        ["14", "POST /tasks missing title",          "400 Bad Request", "✔ PASS"],
        ["15", "GET /tasks with auth",               "200 OK + array",  "✔ PASS"],
        ["16", "PUT /tasks/:id update with auth",    "200 OK",          "✔ PASS"],
        ["17", "DELETE /tasks/:id with auth",        "200 OK",          "✔ PASS"],
        ["18", "DB: password is bcrypt hash ($2...)", "Starts with $2", "✔ PASS"],
    ]
)

doc.add_paragraph()
doc.add_paragraph(
    "[NOTE: Screenshots of the test output, Postman requests, frontend login/register/tasks pages, "
    "and MongoDB Compass showing the bcrypt hash should be inserted here manually by the student.]"
).runs[0].italic = True

doc.add_page_break()

# ══════════════════════════════════════════════════
# 6. API ENDPOINT TABLE
# ══════════════════════════════════════════════════

add_heading_styled("6. API Endpoint Specification", level=1)

add_table(
    ["Method", "Endpoint", "Auth", "Description", "Success", "Error"],
    [
        ["POST",   "/api/auth/register", "No",  "Register new user",       "201", "400, 409"],
        ["POST",   "/api/auth/login",    "No",  "Login, receive JWT",      "200", "400, 401"],
        ["GET",    "/api/auth/me",       "Yes", "Current user profile",    "200", "401, 404"],
        ["GET",    "/tasks",             "Yes", "Get all tasks",           "200", "401, 500"],
        ["GET",    "/tasks/:id",         "Yes", "Get task by ID",          "200", "401, 404"],
        ["POST",   "/tasks",             "Yes", "Create task",             "201", "400, 401"],
        ["PUT",    "/tasks/:id",         "Yes", "Update task",             "200", "400, 401, 404"],
        ["DELETE", "/tasks/:id",         "Yes", "Delete task",             "200", "401, 404"],
    ]
)

doc.add_page_break()

# ══════════════════════════════════════════════════
# 7. KEY QUESTIONS
# ══════════════════════════════════════════════════

add_heading_styled("7. Key Questions", level=1)

add_heading_styled("Q1: Why hash passwords even in a demo project?", level=2)
doc.add_paragraph(
    "Storing plaintext passwords is dangerous regardless of context. If an attacker gains database "
    "access, plaintext passwords are instantly compromised—and since users commonly reuse passwords "
    "across services, a single breach can cascade. bcrypt is specifically designed for password storage: "
    "it incorporates a random salt (defeating precomputed rainbow-table attacks) and a configurable "
    "cost factor (making brute-force attempts computationally expensive). By hashing from day one, "
    "even in a demo, we build security-first habits and ensure no accidental exposure if the database "
    "is ever shared, backed up, or inspected in logs."
)

add_heading_styled("Q2: What does the auth middleware verify, and what happens with a missing or expired token?", level=2)
doc.add_paragraph(
    "The authentication middleware performs three checks in sequence:\n"
    "1. Header presence: It looks for the Authorization header and checks it starts with 'Bearer '. "
    "If absent or malformed, it immediately returns 401.\n"
    "2. Signature verification: jwt.verify() cryptographically checks that the token was signed with "
    "the server's JWT_SECRET. A tampered token (any bit changed) fails verification → 401.\n"
    "3. Expiry check: JWT tokens carry an 'exp' claim. jwt.verify() automatically rejects expired "
    "tokens by throwing a TokenExpiredError, which the middleware catches and returns as a 401 with "
    "a descriptive 'Token has expired' message.\n\n"
    "The entire verification is wrapped in try/catch so that no JWT error can crash the server. "
    "On success, the decoded payload (containing the user's MongoDB _id) is attached to req.user, "
    "allowing downstream controllers to identify the authenticated user."
)

add_heading_styled("Q3: Why is server-side validation needed even if the frontend validates?", level=2)
doc.add_paragraph(
    "Frontend validation exists to improve user experience by providing instant feedback, but it "
    "offers zero security guarantees. Any user can bypass frontend checks entirely by sending direct "
    "HTTP requests via curl, Postman, browser DevTools, or automated scripts. The server is the only "
    "enforceable trust boundary—it must independently validate every input regardless of its source.\n\n"
    "Server-side validation also enables defense in depth: it sanitises inputs (stripping HTML tags "
    "to prevent XSS), enforces business rules (minimum password length, required fields, enum constraints), "
    "and ensures data integrity across all clients (web, mobile, third-party integrations). Without it, "
    "a single crafted request could insert invalid or malicious data directly into the database."
)

doc.add_page_break()

# ══════════════════════════════════════════════════
# 8. ERRORS FACED & REFLECTION
# ══════════════════════════════════════════════════

add_heading_styled("8. Errors Faced and How I Fixed Them", level=1)

errors_faced = [
    (
        "CORS blocked frontend requests to backend",
        "The React dev server (port 5173) and Express backend (port 3000) are on different origins. "
        "Fixed by installing the 'cors' npm package and adding app.use(cors()) in server.js. "
        "Additionally configured Vite's dev proxy to forward /api and /tasks requests to the backend."
    ),
    (
        "Password hash was returned in API responses",
        "Initially, the register and login endpoints included the password hash in the JSON response. "
        "Fixed by adding a toJSON method on the User schema that deletes the password field from the "
        "serialised object, ensuring the hash never leaves the server."
    ),
    (
        "Token verification crashed the server on invalid tokens",
        "When jwt.verify() received a malformed or tampered token, it threw an uncaught exception. "
        "Fixed by wrapping the verification in try/catch and returning a structured 401 response "
        "for each error type (TokenExpiredError vs. general JsonWebTokenError)."
    ),
    (
        "Mongoose duplicate key error (E11000) not handled gracefully",
        "When a user tried to register with an existing email, Mongoose threw an unhandled E11000 error. "
        "Fixed by checking for the existing user before creation (returning 409), and also adding "
        "err.code === 11000 handling in the global error handler as a safety net."
    ),
]

for title, desc in errors_faced:
    p = doc.add_paragraph()
    r = p.add_run(f"Error: {title}")
    r.bold = True
    r.font.size = Pt(10)
    p = doc.add_paragraph(f"Solution: {desc}")
    p.runs[0].font.size = Pt(10)
    doc.add_paragraph()

doc.add_page_break()

# ══════════════════════════════════════════════════
# 9. SUPPLEMENTARY PROBLEMS
# ══════════════════════════════════════════════════

add_heading_styled("9. Supplementary Problems — Implementation", level=1)

supp_items = [
    ("GET /api/auth/me (Protected)",
     "Returns the authenticated user's profile by reading req.user.id (set by the auth middleware) "
     "and querying MongoDB. The password is automatically stripped by the toJSON transform."),
    ("Frontend Register & Login Pages",
     "Two new React pages (Register.jsx, Login.jsx) with controlled form inputs. On successful "
     "registration, the user is redirected to login. On successful login, the JWT is stored in "
     "localStorage and the user is redirected to /tasks."),
    ("Token Storage & Bearer Header",
     "AuthContext stores the token in localStorage and provides an authFetch() wrapper that "
     "automatically attaches the 'Authorization: Bearer <token>' header to every API request."),
    ("401 Handling",
     "The authFetch() wrapper checks every response status. If 401 is received, it clears the "
     "token from localStorage, sets user to null, and React Router navigates to /login."),
    ("Logout Button",
     "The Navbar displays a Logout button when the user is authenticated. Clicking it calls "
     "logout() from AuthContext, which removes the token from localStorage and resets state. "
     "The user is navigated to /login via useNavigate()."),
    ("Hiding Task UI",
     "The /tasks route is wrapped in a ProtectedRoute component that checks for an authenticated user. "
     "If not logged in, it renders <Navigate to='/login' />. The Tasks link in the Navbar is also "
     "conditionally rendered only when user is not null."),
]

for title, desc in supp_items:
    add_heading_styled(title, level=2)
    doc.add_paragraph(desc)

doc.add_page_break()

# ══════════════════════════════════════════════════
# 10. CONCLUSION
# ══════════════════════════════════════════════════

add_heading_styled("10. Conclusion", level=1)
doc.add_paragraph(
    "This practical successfully implements a complete authentication and authorization layer on top "
    "of the existing Task Manager API. Key accomplishments include:\n\n"
    "• User registration with bcrypt password hashing (10 salt rounds)\n"
    "• JWT-based login with 1-hour token expiry\n"
    "• Authentication middleware protecting all task routes with robust error handling\n"
    "• Server-side validation middleware that sanitises inputs and enforces business rules\n"
    "• Frontend integration with React (AuthContext, ProtectedRoute, Login/Register/Tasks pages)\n"
    "• Automatic 401 handling with token clearance and redirect to login\n"
    "• 18 automated tests covering all auth flows and edge cases\n\n"
    "The middleware pipeline follows the order: CORS → JSON Parser → Auth → Validation → Controller → "
    "Error Handler, ensuring that every request is authenticated and validated before reaching "
    "the database. This architecture demonstrates defense-in-depth security principles suitable "
    "for production applications."
)

doc.add_paragraph()
p = doc.add_paragraph()
r = p.add_run("GitHub Repository: ")
r.bold = True
r = p.add_run("https://github.com/RushilNariya10/Task_Manager_API")
r.font.color.rgb = RGBColor(0x0D, 0x6E, 0xFD)

# ── Save ──
doc.save(OUT)
print(f"✔ Lab report saved to: {OUT}")

